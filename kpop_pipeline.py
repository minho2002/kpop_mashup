# -*- coding: utf-8 -*-
"""
K-POP 코드 진행 기반 발매연도 예측 프로젝트 — 공통 파이프라인 모듈.

실험 3~5 노트북에서 확정된 최종 파서 + 피처 엔지니어링 로직을 그대로 옮겨,
Streamlit 앱과 노트북이 동일한 로직을 공유하도록 합니다.
"""
import re
from collections import Counter

import numpy as np
import openpyxl
import pandas as pd

GENDER_MAP = {
    "3세대 걸그룹": ("3세대", "걸그룹"), "3세대 보이그룹": ("3세대", "보이그룹"),
    "3.5세대": ("3.5세대", None), "4세대 걸그룹": ("4세대", "걸그룹"), "4세대 보이그룹": ("4세대", "보이그룹"),
}
DATE_STR_PAT = re.compile(r"^\d{4}[.\-]\d{1,2}[.\-]\d{1,2}$")
DUR_PAT = re.compile(r"^(?P<title>.*?)\s*\((?P<mm>\d+):(?P<ss>\d{2})\)\s*$")
DIATONIC_LABELS = {"1", "2", "3", "4", "5", "6", "7"}
REST_VALUES = {"N.C.", "N.C", "NC", "N.c", "n.c", "-", ""}
NOTE_TO_PC = {
    "C": 0, "C#": 1, "Db": 1, "D": 2, "D#": 3, "Eb": 3, "E": 4, "Fb": 4,
    "E#": 5, "F": 5, "F#": 6, "Gb": 6, "G": 7, "G#": 8, "Ab": 8, "A": 9,
    "A#": 10, "Bb": 10, "B": 11, "Cb": 11, "B#": 0,
}
CHORD_PAT = re.compile(r"^([A-Ga-g])([#bB]?)(m|M)?$")
DEG_BASE_PAT = re.compile(r"b?(\d)")

# 곡 구조(섹션 라벨) 분류용 -- brake/break 오타 변형까지 포함
_BREAK_WORD = r"(?:break|brake)"
RE_BRIDGE = re.compile(r"bridge", re.I)
RE_RAP = re.compile(r"\brap\b", re.I)
RE_BREAK_DANCE = re.compile(rf"(dance[-\s]?{_BREAK_WORD}|{_BREAK_WORD}[-\s]?dance)", re.I)
RE_BREAK_DOWN = re.compile(rf"({_BREAK_WORD}[-\s]?down|drop[-\s]?down|\bdrop\b)", re.I)


def classify_section_label(lbl):
    """섹션 라벨 텍스트 -> {'bridge','rap','break_dance','break_down'} 중 해당하는 것들(리스트)."""
    l = str(lbl).lower()
    types = []
    if RE_BRIDGE.search(l):
        types.append("bridge")
    if RE_RAP.search(l):
        types.append("rap")
    if RE_BREAK_DANCE.search(l):
        types.append("break_dance")
    if RE_BREAK_DOWN.search(l):
        types.append("break_down")
    return types


def normalize_part_label(raw_label):
    """원본 섹션 라벨(예: chorus1, pre-chorus, verse2) -> 매쉬업 유사도 비교용 '파트 카테고리'.
    번호가 붙은 변형(chorus1/2/3, verse1/2 등)은 같은 파트로 합친다. pre-/post-/intro-chorus는
    일반 chorus와 다른 역할을 하므로 별도 카테고리로 유지한다."""
    l = str(raw_label).strip().lower()
    # "chours"(r/u 순서가 뒤바뀐 오타) 등 원본 엑셀에 실제로 존재하는 변형까지 관대하게 매칭
    _CH = r"cho[ru]{2}s"
    if re.search(rf"pre[-\s]?{_CH}", l):
        return "pre-chorus"
    if re.search(rf"post[-\s]?{_CH}", l):
        return "post-chorus"
    if re.search(rf"intro[-\s]?{_CH}", l):
        return "intro-chorus"
    if re.search(_CH, l):
        return "chorus"
    if re.search(r"verse", l):
        return "verse"
    if RE_BRIDGE.search(l):
        return "bridge"
    if re.search(r"build[-\s]?up", l):
        return "build-up"
    if RE_BREAK_DANCE.search(l):
        return "break-dance"
    if RE_BREAK_DOWN.search(l):
        return "break-down"
    if RE_RAP.search(l):
        return "rap"
    if re.search(r"outro", l):
        return "outro"
    if re.search(r"intro", l):
        return "intro"
    return l


# 구간 라벨 셀(예: "intro", "chorus")이 이 색으로 채워져 있으면 그 구간은 2박 기준(마디당 코드 2개)이라는
# 원본 엑셀의 표시 규칙. data_only=True로 열어도 셀 채우기 색은 그대로 유지된다.
TWO_BEAT_FILL_RGB = "FFFFFF00"


def is_two_beat_label_cell(ws, title_row, col):
    """구간 라벨 셀의 배경색으로 2박 기준 여부를 판정한다(노란색=2박 기준)."""
    try:
        cell = ws.cell(row=title_row, column=col)
        rgb = cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor else None
        return isinstance(rgb, str) and rgb.upper() == TWO_BEAT_FILL_RGB
    except Exception:
        return False

MAJOR_LABEL = {
    (0, False): "1", (0, True): "1m", (1, False): "b2", (1, True): "b2m",
    (2, False): "2M", (2, True): "2", (3, False): "b3", (3, True): "b3m",
    (4, False): "3M", (4, True): "3", (5, False): "4", (5, True): "4m",
    (6, False): "b5", (6, True): "b5m", (7, False): "5", (7, True): "5m",
    (8, False): "b6", (8, True): "b6m", (9, False): "6M", (9, True): "6",
    (10, False): "b7", (10, True): "b7m", (11, False): "7", (11, True): "7m",
}
MINOR_LABEL = {
    (0, False): "1M", (0, True): "1", (1, False): "b2", (1, True): "b2m",
    (2, False): "2", (2, True): "2m", (3, False): "3", (3, True): "3m",
    (4, False): "3M2", (4, True): "3m2", (5, False): "4M", (5, True): "4",
    (6, False): "b5", (6, True): "b5m", (7, False): "5M", (7, True): "5",
    (8, False): "6", (8, True): "6m", (9, False): "6M2", (9, True): "6m2",
    (10, False): "7", (10, True): "7m", (11, False): "7M2", (11, True): "7m2",
}
MAJOR_FUNC = {
    (0, False): "T", (0, True): "T", (1, False): "S", (1, True): "S",
    (2, False): "S", (2, True): "S", (3, False): "T", (3, True): "T",
    (4, False): "D", (4, True): "T", (5, False): "S", (5, True): "S",
    (6, False): "D", (6, True): "D", (7, False): "D", (7, True): "D",
    (8, False): "S", (8, True): "S", (9, False): "T", (9, True): "T",
    (10, False): "D", (10, True): "S", (11, False): "D", (11, True): "D",
}
MINOR_FUNC = {
    (0, False): "T", (0, True): "T", (1, False): "S", (1, True): "S",
    (2, False): "S", (2, True): "S", (3, False): "T", (3, True): "T",
    (4, False): "T", (4, True): "T", (5, False): "S", (5, True): "S",
    (6, False): "D", (6, True): "D", (7, False): "D", (7, True): "D",
    (8, False): "T", (8, True): "T", (9, False): "D", (9, True): "T",
    (10, False): "D", (10, True): "D", (11, False): "D", (11, True): "D",
}

NAMED_PROGRESSIONS = {"1564": "1564", "6415": "6415", "1645": "1645", "6451": "6451", "251": "251", "1451": "1451"}

NUMERIC_FEATURES = [
    "duration_sec", "tempo_bpm", "n_sections", "total_measures",
    "borrowed_ratio", "unmatched_ratio", "chord_change_rate",
    "prog_1564", "prog_6415", "prog_1645", "prog_6451", "prog_251", "prog_1451", "prog_any_named_count",
    "bigram_top_ratio", "bigram_entropy", "bigram_unique_count",
    "tsd_T_ratio", "tsd_S_ratio", "tsd_D_ratio", "tsd_DtoT_rate",
    "has_bridge", "has_rap", "has_break_dance", "has_break_down", "n_special_types", "bridge_novel_ratio",
]
# 다중공선성(n_sections<->total_measures, r=0.96) 때문에 실험4에서 채택된 모델링용 최종 피처.
# + 명명된 코드진행 6종(prog_1564/6415/1645/6451/251/1451)과 그 합계(prog_any_named_count)는
#   출현율이 2.7~13%에 불과하고 release_year와의 상관계수가 전부 |r|<0.11, 순열 중요도도 0에 가깝거나
#   음수(노이즈)여서 제외. 실제로 빼도 5-fold Ridge 기준 R²/MAE 변화가 오차범위 수준(0.7021->0.7013)이었음.
_PROG_FEATURES = ["prog_1564", "prog_6415", "prog_1645", "prog_6451", "prog_251", "prog_1451", "prog_any_named_count"]
# + unmatched_ratio는 전 곡에서 항상 0(분산=0, 완전히 죽은 피처), borrowed_ratio는 상관계수(-0.014)와
#   순열 중요도(-0.004) 둘 다 노이즈 수준이어서 제외. tsd_T/S/D_ratio 개별 비율(합쳐서 100%)도
#   상관계수/순열 중요도 모두 낮고 이미 tsd_DtoT_rate가 그 정보를 압축해서 담고 있어 중복이라 제외,
#   tsd_DtoT_rate만 유지.
_LOW_SIGNAL_FEATURES = ["unmatched_ratio", "borrowed_ratio", "tsd_T_ratio", "tsd_S_ratio", "tsd_D_ratio"]
# has_rap/has_break_dance/has_break_down은 n_special_types(네 요소 합)에 이미 포함된 정보라 개별로는
# 중복이어서 모델링 피처에서는 제외하고, has_bridge/n_special_types/bridge_novel_ratio만 사용.
_STRUCT_REDUNDANT = ["has_rap", "has_break_dance", "has_break_down"]
NUMERIC_FEATURES_MODELING = [
    f for f in NUMERIC_FEATURES
    if f != "n_sections" and f not in _PROG_FEATURES and f not in _LOW_SIGNAL_FEATURES and f not in _STRUCT_REDUNDANT
]

FEATURE_KO = {
    "duration_sec": "곡 길이(초)", "tempo_bpm": "템포(BPM)", "n_sections": "섹션 수",
    "total_measures": "전체 마디 수", "borrowed_ratio": "차용화음 비율", "unmatched_ratio": "미분류 코드 비율",
    "chord_change_rate": "코드 전환 빈도", "prog_1564": "1-5-6-4 진행 횟수", "prog_6415": "6-4-1-5 진행 횟수",
    "prog_1645": "1-6-4-5 진행 횟수", "prog_6451": "6-4-5-1 진행 횟수", "prog_251": "2-5-1 진행 횟수",
    "prog_1451": "1-4-5-1 진행 횟수", "prog_any_named_count": "명명된 진행 총합", "bigram_top_ratio": "최빈 2그램 비율",
    "bigram_entropy": "코드 전이 엔트로피", "bigram_unique_count": "고유 2그램 수", "tsd_T_ratio": "Tonic 비율",
    "tsd_S_ratio": "Subdominant 비율", "tsd_D_ratio": "Dominant 비율", "tsd_DtoT_rate": "D→T 해결 비율",
    "has_bridge": "브릿지 유무", "has_rap": "랩 파트 유무", "has_break_dance": "브레이크댄스 유무",
    "has_break_down": "브레이크다운/드롭 유무", "n_special_types": "특수 섹션 종류 수",
    "bridge_novel_ratio": "브릿지 신규코드 비율",
}


def is_date_cell(v):
    if isinstance(v, str) and DATE_STR_PAT.match(v.strip()):
        return True
    if isinstance(v, int) and 19000101 <= v <= 20301231 and len(str(v)) == 8:
        return True
    if hasattr(v, "year"):
        return True
    return False


def parse_date(v):
    try:
        if isinstance(v, str):
            s = v.strip().replace("-", ".")
            parts = [p for p in s.split(".") if p != ""]
            if len(parts) == 3:
                return pd.Timestamp(int(parts[0]), int(parts[1]), int(parts[2]))
            if len(parts) == 2:
                return pd.Timestamp(int(parts[0]), int(parts[1]), 1)
            if len(parts) == 1:
                return pd.Timestamp(int(parts[0]), 1, 1)
        elif isinstance(v, int):
            s = str(v)
            if len(s) == 8:
                return pd.Timestamp(int(s[:4]), int(s[4:6]), int(s[6:8]))
        elif hasattr(v, "year"):
            return pd.Timestamp(v)
    except Exception:
        pass
    return pd.NaT


def parse_title_duration(raw):
    if raw is None:
        return None, None
    s = str(raw).strip()
    m = DUR_PAT.match(s)
    if m:
        return m.group("title").strip(), int(m.group("mm")) * 60 + int(m.group("ss"))
    return s, None


def normalize_chord_text(raw):
    s = str(raw).strip()
    if len(s) <= 1:
        return s
    return s[0] + s[1:].replace("B", "b").replace("M", "m")


def chord_root_quality(chord_text):
    m = CHORD_PAT.match(chord_text)
    if not m:
        return None
    root = m.group(1).upper() + m.group(2)
    pc = NOTE_TO_PC.get(root)
    if pc is None:
        return None
    return pc, (m.group(3) == "m")


def build_lookup_table(ws):
    header = [ws.cell(row=1, column=c).value for c in range(1, 30)]
    major_labels, minor_labels = header[1:17], header[18:29]
    table = {}
    for r in range(2, 14):
        vals = [ws.cell(row=r, column=c).value for c in range(1, 30)]
        if vals[0]:
            degs = {str(l): ch for l, ch in zip(major_labels, vals[1:17]) if l is not None and ch is not None}
            table[str(vals[0])] = {"mode": "major", "degrees": degs}
        if vals[17]:
            degs = {str(l): ch for l, ch in zip(minor_labels, vals[18:29]) if l is not None and ch is not None}
            table[str(vals[17])] = {"mode": "minor", "degrees": degs}
    return table


def compute_degree_and_func(chord_val, key_str):
    if chord_val is None:
        return None, None
    raw = str(chord_val).strip()
    if raw in REST_VALUES:
        return None, "REST"
    norm = normalize_chord_text(raw)
    key_norm = normalize_chord_text(str(key_str).strip()) if key_str else None
    if not key_norm:
        return None, None
    key_is_minor = key_norm.endswith("m")
    key_root_text = key_norm[:-1] if key_is_minor else key_norm
    krq = chord_root_quality(key_root_text)
    cq = chord_root_quality(norm)
    if krq is None or cq is None:
        return None, None
    interval = (cq[0] - krq[0]) % 12
    label_table = MINOR_LABEL if key_is_minor else MAJOR_LABEL
    func_table = MINOR_FUNC if key_is_minor else MAJOR_FUNC
    return label_table.get((interval, cq[1])), func_table.get((interval, cq[1]))


def classify_chord_category(chord_val, key_str, lookup_table):
    if chord_val is None:
        return None
    raw = str(chord_val).strip()
    if raw in REST_VALUES:
        return "REST"
    norm = normalize_chord_text(raw)
    entry = lookup_table.get(str(key_str).strip()) if key_str else None
    if entry is not None:
        for label, ch in entry["degrees"].items():
            if str(ch).strip() == norm:
                return "diatonic" if label in DIATONIC_LABELS else "borrowed"
        cq = chord_root_quality(norm)
        if cq is not None:
            for label, ch in entry["degrees"].items():
                cand = chord_root_quality(normalize_chord_text(str(ch).strip()))
                if cand is not None and cand == cq:
                    return "diatonic" if label in DIATONIC_LABELS else "borrowed"
    label, func = compute_degree_and_func(chord_val, key_str)
    if label is not None:
        return "diatonic" if label in DIATONIC_LABELS else "borrowed"
    return "unmatched"


def base_degree(label):
    if label is None:
        return None
    m = DEG_BASE_PAT.match(str(label))
    return m.group(1) if m else None


def parse_sheet(ws, sheet_name):
    max_col, max_row = ws.max_column, ws.max_row
    date_rows = [r for r in range(1, max_row + 1) if is_date_cell(ws.cell(row=r, column=2).value)]
    last_artist, artist_by_row = None, {}
    for r in range(1, max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v not in (None, ""):
            last_artist = str(v).strip()
        artist_by_row[r] = last_artist

    lookup = build_lookup_table(ws)
    songs = []

    for date_row in date_rows:
        title_row, chord_row, degree_row = date_row - 3, date_row - 2, date_row - 1
        if title_row < 1:
            continue
        gen, gender = GENDER_MAP.get(sheet_name, (sheet_name, None))
        artist = artist_by_row.get(title_row)
        title, duration_sec = parse_title_duration(ws.cell(row=title_row, column=2).value)
        key_val = ws.cell(row=chord_row, column=2).value
        tempo_val = ws.cell(row=degree_row, column=2).value
        date_parsed = parse_date(ws.cell(row=date_row, column=2).value)

        label_events = [
            (c, str(ws.cell(row=title_row, column=c).value).strip())
            for c in range(3, max_col + 1)
            if ws.cell(row=title_row, column=c).value not in (None, "")
        ]
        last_used_col = 2
        for c in range(3, max_col + 1):
            if ws.cell(row=chord_row, column=c).value not in (None, ""):
                last_used_col = c
        if last_used_col <= 2 and label_events:
            last_used_col = label_events[-1][0]

        has_chords = any(ws.cell(row=chord_row, column=c).value not in (None, "") for c in range(3, max_col + 1))
        n_labels = len(label_events)

        all_chords, all_categories, all_degrees, all_tsd = [], [], [], ""
        section_types_present, bridge_chords, other_chords = set(), set(), set()
        sections_list = []
        for i, (col, raw_label) in enumerate(label_events):
            col_end = (label_events[i + 1][0] - 1) if i + 1 < n_labels else last_used_col
            if col_end < col:
                continue
            label_types = classify_section_label(raw_label)
            section_types_present.update(label_types)
            chords_here = set()
            sec_chords, sec_degrees, sec_tsd = [], [], ""
            for c in range(col, col_end + 1):
                cv = ws.cell(row=chord_row, column=c).value
                if cv not in (None, ""):
                    cv_s = str(cv).strip()
                    all_chords.append(cv_s)
                    all_categories.append(classify_chord_category(cv, key_val, lookup))
                    label, func = compute_degree_and_func(cv, key_val)
                    deg_label = label if label else ""
                    func_char = "R" if func == "REST" else (func if func else "?")
                    all_degrees.append(deg_label)
                    all_tsd += func_char
                    sec_chords.append(cv_s)
                    sec_degrees.append(deg_label)
                    sec_tsd += func_char
                    if cv_s not in REST_VALUES:
                        chords_here.add(cv_s)
            sections_list.append({
                "raw_label": raw_label, "part": normalize_part_label(raw_label),
                "chords": sec_chords, "degrees": sec_degrees, "tsd": sec_tsd,
                "n_chords": len(sec_chords), "two_beat": is_two_beat_label_cell(ws, title_row, col),
            })
            if "bridge" in label_types:
                bridge_chords |= chords_here
            else:
                other_chords |= chords_here

        has_bridge = "bridge" in section_types_present
        has_rap = "rap" in section_types_present
        has_break_dance = "break_dance" in section_types_present
        has_break_down = "break_down" in section_types_present
        n_special_types = sum([has_bridge, has_rap, has_break_dance, has_break_down])
        if bridge_chords:
            bridge_novel_ratio = len(bridge_chords - other_chords) / len(bridge_chords)
        else:
            bridge_novel_ratio = np.nan

        songs.append({
            "sheet": sheet_name, "generation": gen, "gender_group": gender, "artist": artist, "title": title,
            "duration_sec": duration_sec, "key": key_val,
            "key_mode": ("minor" if isinstance(key_val, str) and key_val.strip().endswith("m") else ("major" if key_val else None)),
            "tempo_bpm": tempo_val, "release_date": date_parsed, "has_chords": has_chords, "n_sections": n_labels,
            "has_bridge": int(has_bridge), "has_rap": int(has_rap),
            "has_break_dance": int(has_break_dance), "has_break_down": int(has_break_down),
            "n_special_types": n_special_types, "bridge_novel_ratio": bridge_novel_ratio,
            "chords": all_chords, "chord_categories": all_categories, "degrees": all_degrees, "tsd": all_tsd,
            "sections": sections_list,
        })
    return pd.DataFrame(songs)


def song_progression_features(degrees):
    bases = [base_degree(d) for d in degrees]
    seq = "".join(b for b in bases if b is not None)
    feats = {}
    for name, pat in NAMED_PROGRESSIONS.items():
        cnt, start = 0, 0
        while True:
            idx = seq.find(pat, start)
            if idx == -1:
                break
            cnt += 1
            start = idx + 1
        feats[f"prog_{name}"] = cnt
    feats["prog_any_named_count"] = sum(feats.values())
    bigrams = [seq[i:i + 2] for i in range(len(seq) - 1)]
    if bigrams:
        c = Counter(bigrams)
        total = len(bigrams)
        top_ratio = c.most_common(1)[0][1] / total
        probs = np.array(list(c.values())) / total
        entropy = float(-(probs * np.log2(probs)).sum())
        n_unique = len(c)
    else:
        top_ratio, entropy, n_unique = np.nan, np.nan, 0
    feats["bigram_top_ratio"] = top_ratio
    feats["bigram_entropy"] = entropy
    feats["bigram_unique_count"] = n_unique
    return feats


def song_tsd_features(tsd):
    tsd = tsd.replace("?", "").replace("R", "")
    if not tsd:
        return {"tsd_T_ratio": np.nan, "tsd_S_ratio": np.nan, "tsd_D_ratio": np.nan, "tsd_DtoT_rate": np.nan}
    n = len(tsd)
    t, s, d = tsd.count("T"), tsd.count("S"), tsd.count("D")
    dt_trans = sum(1 for i in range(n - 1) if tsd[i] == "D" and tsd[i + 1] == "T")
    d_non_last = sum(1 for i in range(n - 1) if tsd[i] == "D")
    return {"tsd_T_ratio": t / n, "tsd_S_ratio": s / n, "tsd_D_ratio": d / n,
            "tsd_DtoT_rate": (dt_trans / d_non_last) if d_non_last else np.nan}


def song_chord_agg(chords, categories):
    total = borrowed = unmatched = 0
    for ch, cat in zip(chords, categories):
        total += 1
        if cat == "borrowed":
            borrowed += 1
        elif cat == "unmatched":
            unmatched += 1
    changes = sum(1 for i in range(1, len(chords)) if chords[i] != chords[i - 1])
    change_rate = changes / (len(chords) - 1) if len(chords) > 1 else np.nan
    return pd.Series({
        "total_measures": total,
        "borrowed_ratio": borrowed / total if total else np.nan,
        "unmatched_ratio": unmatched / total if total else np.nan,
        "chord_change_rate": change_rate,
    })


def load_and_engineer(xlsx_path):
    """KPOP MASHUP.xlsx -> (songs_df, modeling_df) 전체 파이프라인 실행."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    all_songs = [parse_sheet(wb[name], name) for name in wb.sheetnames]
    songs_df = pd.concat(all_songs, ignore_index=True)

    modeling_df = songs_df[songs_df["has_chords"]].copy()
    modeling_df["release_year"] = modeling_df["release_date"].dt.year
    modeling_df["duration_sec"] = modeling_df["duration_sec"].fillna(modeling_df["duration_sec"].median())
    modeling_df["tempo_bpm"] = modeling_df["tempo_bpm"].fillna(modeling_df["tempo_bpm"].median())
    for col in ["key_mode", "generation", "gender_group"]:
        modeling_df[col] = modeling_df[col].astype("category")

    chord_feats = modeling_df.apply(lambda r: song_chord_agg(r["chords"], r["chord_categories"]), axis=1)
    prog_feats = modeling_df["degrees"].apply(song_progression_features).apply(pd.Series)
    tsd_feats = modeling_df["tsd"].apply(song_tsd_features).apply(pd.Series)
    modeling_df = pd.concat([modeling_df.reset_index(drop=True), chord_feats.reset_index(drop=True),
                              prog_feats.reset_index(drop=True), tsd_feats.reset_index(drop=True)], axis=1)
    modeling_df["tsd_DtoT_rate"] = modeling_df["tsd_DtoT_rate"].fillna(0.0)
    # 브릿지 없는 곡은 "브릿지발 신규 코드"도 없다는 뜻이므로 0으로 채움 (NaN이면 dropna에서 곡 전체가 빠짐)
    modeling_df["bridge_novel_ratio"] = modeling_df["bridge_novel_ratio"].fillna(0.0)

    modeling_df = modeling_df.dropna(subset=["release_year"] + NUMERIC_FEATURES).reset_index(drop=True)
    return songs_df, modeling_df


def build_model_matrix(modeling_df, features):
    """더미 인코딩된 전체 X, y, 컬럼명을 반환. generation/gender_group은 원본 라벨도 별도 컬럼으로 보존."""
    df = modeling_df.dropna(subset=["release_year"] + features).copy()
    df["generation_label"] = df["generation"].astype(str)
    df["gender_group_label"] = df["gender_group"].astype(str)
    df = pd.get_dummies(df, columns=["key_mode", "generation", "gender_group"], prefix=["mode", "gen", "grp"])
    df = df.reset_index(drop=True)
    # 베이스라인(major/3세대/걸그룹)은 드롭하고 나머지는 어떤 세대 카테고리가 새로 생기든(예: 2.5세대) 자동 포함
    drop_baseline = {"mode_major", "gen_3세대", "grp_걸그룹"}
    dummy_cols = [c for c in df.columns if c.startswith(("mode_", "gen_", "grp_")) and c not in drop_baseline]
    return df, features + dummy_cols, dummy_cols


def get_models():
    from sklearn.linear_model import LinearRegression, Ridge, Lasso
    from sklearn.ensemble import RandomForestRegressor

    class MeanBaseline:
        def fit(self, X, y):
            self.mean_ = float(np.mean(y))
            return self

        def predict(self, X):
            return np.full(len(X), self.mean_)

    return {
        "평균 베이스라인": MeanBaseline,
        "Linear Regression": LinearRegression,
        "Ridge (alpha=1.0)": lambda: Ridge(alpha=1.0, random_state=42),
        "Lasso (alpha=0.1)": lambda: Lasso(alpha=0.1, random_state=42),
        "Random Forest": lambda: RandomForestRegressor(n_estimators=300, max_depth=6, min_samples_leaf=3, random_state=42),
    }


def run_time_split_eval(modeling_df, features, split_cutoff=None, test_ratio=0.15):
    """시간 기준 분할 평가. split_cutoff=None이면 뒤 test_ratio 비율을 test로 사용, 문자열('YYYY-MM-DD')이면 날짜 컷."""
    from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score

    df, cols, _ = build_model_matrix(modeling_df, features)
    df = df.sort_values("release_date").reset_index(drop=True)

    if split_cutoff is None:
        n_test = max(1, int(len(df) * test_ratio))
        train_df, test_df = df.iloc[:-n_test].copy(), df.iloc[-n_test:].copy()
    else:
        train_df = df[df["release_date"] < split_cutoff].copy()
        test_df = df[df["release_date"] >= split_cutoff].copy()

    means = train_df[features].mean()
    stds = train_df[features].std().replace(0, 1)
    for col in features:
        train_df[col + "_s"] = (train_df[col] - means[col]) / stds[col]
        test_df[col + "_s"] = (test_df[col] - means[col]) / stds[col]
    dummy_cols = [c for c in cols if c not in features]
    scaled_cols = [f + "_s" for f in features] + dummy_cols

    X_train, y_train = train_df[scaled_cols].astype(float), train_df["release_year"].astype(float)
    X_test, y_test = test_df[scaled_cols].astype(float), test_df["release_year"].astype(float)

    rows, preds = [], {}
    for name, build in get_models().items():
        model = build()
        model.fit(X_train, y_train)
        pred = model.predict(X_test)
        preds[name] = pred
        rows.append({
            "model": name, "MAE": mean_absolute_error(y_test, pred),
            "RMSE": float(np.sqrt(mean_squared_error(y_test, pred))), "R2": r2_score(y_test, pred),
        })
    result_df = pd.DataFrame(rows).sort_values("RMSE").reset_index(drop=True)
    return result_df, test_df, y_test.values, preds, len(train_df), len(test_df)


def run_kfold_eval(modeling_df, features, k=5, seed=42):
    from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
    from sklearn.model_selection import KFold

    df, cols, dummy_cols = build_model_matrix(modeling_df, features)
    kf = KFold(n_splits=k, shuffle=True, random_state=seed)
    models = get_models()
    oof_pred = {name: np.zeros(len(df)) for name in models}
    y_all = df["release_year"].astype(float).values

    for train_idx, test_idx in kf.split(df):
        train_df, test_df = df.iloc[train_idx].copy(), df.iloc[test_idx].copy()
        means = train_df[features].mean()
        stds = train_df[features].std().replace(0, 1)
        for col in features:
            train_df[col + "_s"] = (train_df[col] - means[col]) / stds[col]
            test_df[col + "_s"] = (test_df[col] - means[col]) / stds[col]
        scaled_cols = [f + "_s" for f in features] + dummy_cols
        X_train, y_train = train_df[scaled_cols].astype(float), train_df["release_year"].astype(float)
        X_test = test_df[scaled_cols].astype(float)
        for name, build in models.items():
            model = build()
            model.fit(X_train, y_train)
            oof_pred[name][test_idx] = model.predict(X_test)

    rows = []
    for name, pred in oof_pred.items():
        rows.append({
            "model": name, "MAE": mean_absolute_error(y_all, pred),
            "RMSE": float(np.sqrt(mean_squared_error(y_all, pred))), "R2": r2_score(y_all, pred),
        })
    result_df = pd.DataFrame(rows).sort_values("RMSE").reset_index(drop=True)
    return result_df, df, y_all, oof_pred


def run_permutation_importance(modeling_df, features, model_name="Ridge (alpha=1.0)", k=5, seed=42, n_repeats=15):
    """K-fold CV 기반 순열 중요도: 각 피처를 fold별 test셋에서 섞었을 때 R2가 얼마나 떨어지는지."""
    from sklearn.metrics import r2_score
    from sklearn.model_selection import KFold

    df, cols, dummy_cols = build_model_matrix(modeling_df, features)
    kf = KFold(n_splits=k, shuffle=True, random_state=seed)
    models = get_models()
    y_all = df["release_year"].astype(float).values

    base_pred = np.zeros(len(df))
    fold_info = []
    for train_idx, test_idx in kf.split(df):
        train_df, test_df = df.iloc[train_idx].copy(), df.iloc[test_idx].copy()
        means = train_df[features].mean()
        stds = train_df[features].std().replace(0, 1)
        for col in features:
            train_df[col + "_s"] = (train_df[col] - means[col]) / stds[col]
            test_df[col + "_s"] = (test_df[col] - means[col]) / stds[col]
        scaled_cols = [f + "_s" for f in features] + dummy_cols
        X_train, y_train = train_df[scaled_cols].astype(float), train_df["release_year"].astype(float)
        model = models[model_name]()
        model.fit(X_train, y_train)
        X_test = test_df[scaled_cols].astype(float)
        base_pred[test_idx] = model.predict(X_test)
        fold_info.append((test_idx, means, stds, model, scaled_cols))

    base_r2 = r2_score(y_all, base_pred)
    rng = np.random.RandomState(123)
    importance = {f: [] for f in features}
    for _ in range(n_repeats):
        for f in features:
            pred_shuf = base_pred.copy()
            for test_idx, means, stds, model, scaled_cols in fold_info:
                test_df = df.iloc[test_idx].copy()
                shuffled = test_df[f].values.copy()
                rng.shuffle(shuffled)
                test_df[f] = shuffled
                for col in features:
                    test_df[col + "_s"] = (test_df[col] - means[col]) / stds[col]
                pred_shuf[test_idx] = model.predict(test_df[scaled_cols].astype(float))
            importance[f].append(base_r2 - r2_score(y_all, pred_shuf))

    imp_df = pd.DataFrame([{"feature": f, "feature_ko": FEATURE_KO.get(f, f), "r2_drop": float(np.mean(v))}
                            for f, v in importance.items()]).sort_values("r2_drop", ascending=False)
    return imp_df, base_r2


def compute_residuals(modeling_df, features, model_name="Ridge (alpha=1.0)", k=5, seed=42):
    """K-fold CV OOF 예측 + 잔차 + (선형모델인 경우) 피처별 기여도 분해."""
    from sklearn.model_selection import KFold

    df, cols, dummy_cols = build_model_matrix(modeling_df, features)
    kf = KFold(n_splits=k, shuffle=True, random_state=seed)
    models = get_models()
    y_all = df["release_year"].astype(float).values
    pred = np.zeros(len(df))
    contrib = np.zeros((len(df), len(features)))
    is_linear = model_name in ("Linear Regression", "Ridge (alpha=1.0)", "Lasso (alpha=0.1)")

    for train_idx, test_idx in kf.split(df):
        train_df, test_df = df.iloc[train_idx].copy(), df.iloc[test_idx].copy()
        means = train_df[features].mean()
        stds = train_df[features].std().replace(0, 1)
        for col in features:
            train_df[col + "_s"] = (train_df[col] - means[col]) / stds[col]
            test_df[col + "_s"] = (test_df[col] - means[col]) / stds[col]
        scaled_cols = [f + "_s" for f in features] + dummy_cols
        X_train, y_train = train_df[scaled_cols].astype(float), train_df["release_year"].astype(float)
        model = models[model_name]()
        model.fit(X_train, y_train)
        X_test = test_df[scaled_cols].astype(float)
        pred[test_idx] = model.predict(X_test)
        if is_linear:
            coef = pd.Series(model.coef_, index=scaled_cols)
            feat_scaled_cols = [f + "_s" for f in features]
            contrib[test_idx, :] = X_test[feat_scaled_cols].values * coef[feat_scaled_cols].values

    out = df[["artist", "title", "release_year", "generation_label", "gender_group_label"]].copy()
    out = out.rename(columns={"generation_label": "generation", "gender_group_label": "gender_group"})
    out["pred_year"] = pred
    out["residual"] = y_all - pred
    out["abs_residual"] = out["residual"].abs()
    out["orig_pos"] = np.arange(len(out))  # contrib 배열의 행 위치 — 이후 정렬/reset_index와 무관하게 곡별 기여도를 정확히 찾기 위함
    return out.sort_values("abs_residual", ascending=False), contrib, features, is_linear
